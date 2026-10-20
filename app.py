import streamlit as st
from google import genai
import docx
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
import io
import re
import time

# ==========================================
# 1. إعدادات الواجهة والصفحة
# ==========================================
st.set_page_config(
    page_title="نظام تحليل محادثات الواتساب - جهاز الأمن الوطني",
    page_icon="📊",
    layout="wide"
)

st.title("📊 نظام تحليل الرأي العام والتقارير الأسبوعية")
st.subheader("دائرة الأمن الوقائي - مديرية العمليات النفسية")

# إدخال مفتاح API وشريط البحث في الشريط الجانبي
with st.sidebar:
    st.header("⚙️ إعدادات النظام")
    api_key = st.text_input(
        "أدخل مفتاح Gemini API الخاص بك:", 
        type="password",
        help="احصل على مفتاح مجاني من https://aistudio.google.com/app/apikey"
    )
    
    st.markdown("---")
    st.header("🔍 البحث عن إحصائيات عضو")
    search_user = st.text_input(
        "أدخل اسم الشخص للبحث عن مشاركاته:",
        placeholder="اكتب أي جزء من الاسم..."
    )
    st.info("💡 يمكنك كتابة جزء من الاسم لمعرفة عدد مرفوعاته داخل المحادثة.")

# ==========================================
# 2. دالة تنظيف وقراءة المحادثات من ملف Word (DOCX)
# ==========================================
def parse_whatsapp_chat_docx(uploaded_file):
    # تحميل المستند من الذاكرة
    doc = docx.Document(io.BytesIO(uploaded_file.getvalue()))
    
    # استخراج النصوص من جميع الفقرات والجداول داخل ملف الـ Word
    lines = []
    for p in doc.paragraphs:
        if p.text.strip():
            lines.append(p.text.strip())
            
    # قراءة النصوص إذا كانت داخل جداول في ملف الـ Word
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    lines.append(cell.text.strip())

    messages = []
    user_stats = {}
    
    for line_str in lines:
        if not line_str or "<Media omitted>" in line_str or "رسائل ومكالمات مشفرة" in line_str or "end-to-end encrypted" in line_str:
            continue
        
        if ":" in line_str:
            parts = line_str.split(":", 1)
            left_part = parts[0]
            msg_text = parts[1].strip() if len(parts) > 1 else ""
            
            # استخراج اسم المرسل بإزالة التاريخ والتوقيت والرموز الخاصة الشائعة
            sender_name = left_part
            if "-" in left_part:
                sender_name = left_part.split("-")[-1]
            elif "]" in left_part:
                sender_name = left_part.split("]")[-1]
            elif ")" in left_part:
                sender_name = left_part.split(")")[-1]
                
            sender_name = sender_name.strip()

            if msg_text:
                messages.append(f"{sender_name}: {msg_text}")
                if sender_name and not sender_name.isdigit():
                    user_stats[sender_name] = user_stats.get(sender_name, 0) + 1
        else:
            if messages:
                messages[-1] += " " + line_str
                
    return messages, user_stats

def get_balanced_sample(msgs, max_count=1000):
    if len(msgs) <= max_count:
        return msgs
    step = len(msgs) // max_count
    return msgs[::step][:max_count]

# ==========================================
# 3. دالة إنشاء مستند Word (DOCX) وتنسيقه
# ==========================================
def create_docx_report(report_text):
    doc = docx.Document()
    lines = report_text.splitlines()
    
    for line in lines:
        line_str = line.strip()
        if not line_str:
            continue
            
        if line_str.startswith("# "):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            run = p.add_run(line_str.replace("# ", ""))
            run.font.size = Pt(18)
            run.font.bold = True
            run.font.color.rgb = RGBColor(0, 51, 102)
        elif line_str.startswith("## "):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            run = p.add_run(line_str.replace("## ", ""))
            run.font.size = Pt(16)
            run.font.bold = True
            run.font.color.rgb = RGBColor(0, 102, 153)
        elif line_str.startswith("### "):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            run = p.add_run(line_str.replace("### ", ""))
            run.font.size = Pt(14)
            run.font.bold = True
        else:
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            run = p.add_run(line_str.replace("**", ""))
            run.font.size = Pt(12)
            
    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer

# ==========================================
# 4. دالة التحليل وبناء التقرير عبر Gemini (مع معالجة خطأ 503)
# ==========================================
def generate_single_chat_report(chat_msgs, user_api_key):
    client = genai.Client(api_key=user_api_key.strip())
    
    sample_msgs = get_balanced_sample(chat_msgs, 1000)
    combined_text = "\n".join([f"- {m}" for m in sample_msgs])

    prompt = f"""
أنت خبير إستراتيجي في تحليل العمليات النفسية والرأي العام لجهاز أمني رسمي.
لديك البيانات والمشاركات المستخرجة من محادثة الجروب المرفق:

---
{combined_text}
---

المطلوب منك تحليل كافة المحادثات والمنشورات بالكامل، وتوليد التقرير النهائي باللغة العربية مع الالتزام بالواجهة والهيكلية أدناه بالضبط:

==================================================
جهاز الأمن الوطني
مديرية أمن بغداد
قسم الوقائي
العمليات النفسية

التقرير الأسبوعي
==================================================

قائمة المحتويات:
1. الإيجاز العام
2. موقف العمليات النفسية وأبرز الشائعات والرأي العام والمرصد
3. تحليل شامل لكافة المواضيع المرصودة بدون استثناء
4. موقف الخطاب الديني والاجتماعي
5. القراءة الاستشرافية والتهديدات المحتملة
6. المعالجات والتوصيات المقترحة
7. جدول ملخص كافة المواضيع وتصنيفاتها

--------------------------------------------------

### 1. الإيجاز العام:
ملخص تنفيذي شامل لكافة ما تم رصده في المحادثة لهذا الأسبوع.

### 2. موقف العمليات النفسية وأبرز الشائعات والرأي العام والمرصد:
تحليل دقيق لكافة تحركات التضليل الإعلامي والحملات الموجهة والشائعات المنتشرة.

### 3. تحليل شامل لكافة المواضيع المرصودة (شامل لجميع المواضيع الفرعية والرئيسية المذكورة في الرسائل):
قم باستخراج وتحليل **جميع** المواضيع المطروحة في المحادثة دون حذف أو اختصار لمواضيع معينة، ولكل موضوع اكتب:
- **عنوان الموضوع**:
- **شرح وتفاصيل الموضوع**:
- **المنصة / المصدر**:
- **الحلول المقترحة**:
- **الاستنتاجات**:
- **التوصيات**:

### 4. موقف الخطاب الديني والاجتماعي:
رصد التوجهات في الخطاب الديني أو العشائري أو الاجتماعي.

### 5. القراءة الاستشرافية والتهديدات المحتملة:
توقع التهديدات والسيناريوهات القادمة بناءً على كل ما تم رصده.

### 6. المعالجات والتوصيات المقترحة:
التوصيات التنفيذية المقترحة للتعامل مع الشائعات وإدارة الرأي العام.

### 7. جدول ملخص كافة المواضيع وتصنيفاتها:
أنشئ جدول Markdown شامل يحتوي على **كافة المواضيع المذكورة في المحادثات**:
| الرقم | الموضوع | نوع التصنيف (شائعة / رأي عام / رصد) | التقييم (إيجابي / سلبي / محايد) | المنصة / المصدر | مستوى التفاعل | المعالجة المقترحة |
"""

    # استخدام النماذج الرسمية المستقرة للتنقل بينها عند وجود ضغط
    models_to_try = ['gemini-1.5-flash', 'gemini-1.5-pro']
    
    max_retries = 3
    for attempt in range(max_retries):
        for model_name in models_to_try:
            try:
                response = client.models.generate_content(
                    model=model_name,
                    contents=prompt
                )
                return response.text
            except Exception as e:
                # إذا كان الخطأ بسبب الضغط المؤقت على السيرفرات (503)، يتم الانتظار ثم المحاولة مجدداً
                if "503" in str(e) or "UNAVAILABLE" in str(e):
                    time.sleep(2 * (attempt + 1))
                    continue
                else:
                    raise e
                    
    raise Exception("سيرفرات جوجل مجهدة حالياً (503)، يرجى إعادة الضغط على زر التوليد بعد بضع ثوانٍ.")

# ==========================================
# 5. واجهة رفع ملف الجروب (بصيغة ورد)
# ==========================================
st.markdown("### 💬 رفع ملف محادثة الجروب بصيغة ورد (.docx)")
single_file = st.file_uploader("يرجى رفع ملف المحادثة الخاص بالجروب", type=["docx", "doc"])

st.markdown("---")

# ==========================================
# 6. زر المعالجة والإحصائيات وتوليد النتائج
# ==========================================
if single_file:
    try:
        msgs, user_stats = parse_whatsapp_chat_docx(single_file)
        total_messages = len(msgs)
        
        st.success(f"✅ تم قراءة ملف الورد بنجاح! إجمالي الرسائل المقروءة: {total_messages} رسالة.")
        
        # عرض الإحصائيات عند البحث
        if search_user and search_user.strip():
            query = search_user.strip().lower()
            matched_users = [u for u in user_stats.keys() if query in u.lower()]
            
            if matched_users:
                st.markdown("### 📈 نتائج البحث عن الشخص:")
                for u in matched_users:
                    count = user_stats[u]
                    percentage = (count / total_messages * 100) if total_messages > 0 else 0
                    st.info(f"👤 **الاسم المرصود**: {u} | 📊 **عدد الرسائل/المرفوعات**: {count} رسالة | 🎯 **نسبة التفاعل**: {percentage:.1f}%")
            else:
                st.warning(f"⚠️ لم يتم العثور على اسم يحتوي على '{search_user}' داخل هذه المحادثة.")

        st.markdown("---")

        if st.button("🚀 توليد التقرير الأسبوعي الشامل", type="primary", use_container_width=True):
            if not api_key:
                st.error("⚠️ يرجى إدخال مفتاح API الخاص بك في الشريط الجانبي أولاً.")
            else:
                try:
                    with st.spinner("⏳ جاري تحليل المحادثة وصياغة التقرير الإستراتيجي..."):
                        report = generate_single_chat_report(msgs, api_key)
                        docx_file = create_docx_report(report)
                        
                        st.markdown("---")
                        st.header("📊 التقرير الأسبوعي المدمج والتحليل الشامل")
                        st.markdown(report)
                        
                        st.download_button(
                            label="📄 تنزيل التقرير بملف ورد (Weekly_Report.docx)",
                            data=docx_file,
                            file_name="Weekly_Report.docx",
                            mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                            use_container_width=True
                        )
                except Exception as e:
                    st.error(f"❌ حدث خطأ أثناء تنفيذ الطلب: {e}")
    except Exception as e:
        st.error(f"❌ حدث خطأ أثناء قراءة ملف الورد: {e}")
